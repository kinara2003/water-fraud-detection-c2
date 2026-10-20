"""
tests/test_slide_generator.py
Unit and Integration Tests for scripts/generate_slides.py.

Verifies:
1. Generation of the 16:9 widescreen presentation deck.
2. Presentation contains exactly 12 slides matching the oral viva defense sequence.
3. Key figures from reports/figures/ are properly embedded into the designated slides.
4. Slide metadata, headers, titles, and layout containers meet visual quality standards.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from scripts.generate_slides import (
    generate_presentation,
    DEFAULT_OUTPUT_PPTX,
    FIGURES_DIR,
)


@pytest.fixture(scope="module")
def generated_deck(tmp_path_factory) -> Path:
    """Generates the viva presentation into a temporary test location once per module."""
    tmp_dir = tmp_path_factory.mktemp("viva_slides")
    out_path = tmp_dir / "Component2_Viva_Test.pptx"
    generated_path = generate_presentation(output_path=out_path)
    return generated_path


def test_presentation_file_generated(generated_deck: Path):
    """Asserts that the presentation file exists, is non-empty, and exceeds 500 KB."""
    assert generated_deck.exists(), "Presentation file was not created"
    file_size = generated_deck.stat().st_size
    assert file_size > 500_000, f"File size ({file_size} bytes) is suspiciously small (expected > 500 KB)"


def test_widescreen_aspect_ratio_and_slide_count(generated_deck: Path):
    """Asserts 16:9 widescreen dimensions and exactly 12 slides."""
    prs = Presentation(str(generated_deck))
    assert len(prs.slides) == 12, f"Expected 12 slides, found {len(prs.slides)}"

    # 16:9 widescreen: 13.333" x 7.5"
    expected_width = Inches(13.333)
    expected_height = Inches(7.5)

    assert abs(prs.slide_width - expected_width) < 1000, f"Slide width mismatch: {prs.slide_width}"
    assert abs(prs.slide_height - expected_height) < 1000, f"Slide height mismatch: {prs.slide_height}"


def test_slide_titles_and_topics(generated_deck: Path):
    """Asserts that each of the 12 slides contains expected keywords in text shapes."""
    prs = Presentation(str(generated_deck))

    expected_keywords = [
        (0, ["Water Theft", "Surveillance", "NWSDB"]),                          # Slide 1: Title & Context
        (1, ["Problem Statement", "Intermittent Water Supply"]),                 # Slide 2: IWS NRW Challenge
        (2, ["Digital Twin", "EPANET", "WNTR"]),                                 # Slide 3: EPANET/WNTR PDD
        (3, ["Feature Engineering", "Hydraulic Deficit"]),                       # Slide 4: Feature Engineering
        (4, ["AI Architecture", "Isolation Forest", "Random Forest"]),          # Slide 5: AI Architecture
        (5, ["Fraud Risk Engine", "IWS", "Transient Filter"]),                   # Slide 6: Risk Engine & Filter
        (6, ["Confusion Matrix", "Quantitative Benchmark"]),                     # Slide 7: Benchmark Results
        (7, ["Cross-Component", "Integration", "C1", "C3", "C4"]),              # Slide 8: C1, C3, C4 Interlock
        (8, ["Dashboard", "Inspection Dispatch"]),                               # Slide 9: Operator Dashboard
        (9, ["Resilience", "Edge Gateway"]),                                     # Slide 10: Edge Resilience
        (10, ["Boundaries", "Ethics", "Enhancements"]),                          # Slide 11: Limitations & Ethics
        (11, ["Live Demonstration", "Conclusion"]),                              # Slide 12: Live Demo Script
    ]

    for slide_idx, keywords in expected_keywords:
        slide = prs.slides[slide_idx]
        all_text = " ".join(
            shape.text_frame.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and shape.text_frame
        )

        for kw in keywords:
            assert kw.lower() in all_text.lower(), (
                f"Slide {slide_idx + 1} missing keyword '{kw}'. Found text snippet: {all_text[:200]}"
            )


def test_embedded_figures_in_designated_slides(generated_deck: Path):
    """Asserts that publication figures from reports/figures/ are embedded into slides 4, 6, 7, and 9."""
    prs = Presentation(str(generated_deck))

    def count_pictures(slide) -> int:
        return sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)

    # Slide 4: Feature Importance Plot
    pics_slide_4 = count_pictures(prs.slides[3])
    assert pics_slide_4 >= 1, f"Slide 4 should contain feature importance figure, found {pics_slide_4}"

    # Slide 6: Refill vs Theft Plot
    pics_slide_6 = count_pictures(prs.slides[5])
    assert pics_slide_6 >= 1, f"Slide 6 should contain refill vs theft trace figure, found {pics_slide_6}"

    # Slide 7: Model Comparison Benchmark & Confusion Matrix
    pics_slide_7 = count_pictures(prs.slides[6])
    assert pics_slide_7 >= 2, f"Slide 7 should contain at least 2 figures (Benchmark & Confusion Matrix), found {pics_slide_7}"

    # Slide 9: Spatial DMA Comparison Plot
    pics_slide_9 = count_pictures(prs.slides[8])
    assert pics_slide_9 >= 1, f"Slide 9 should contain spatial delta DMA figure, found {pics_slide_9}"


def test_default_output_pptx_exists():
    """Asserts that reports/Component2_Viva_Presentation.pptx exists and is non-empty."""
    assert DEFAULT_OUTPUT_PPTX.exists(), (
        f"Default presentation not found at {DEFAULT_OUTPUT_PPTX}. Run scripts/generate_slides.py first."
    )
    assert DEFAULT_OUTPUT_PPTX.stat().st_size > 1_000_000, "Default presentation size should exceed 1 MB"
